#!/usr/bin/env python3
"""Export the dashboard's MART tables as the BI user, and check the export (Phase 6).

    python scripts/export_marts.py                     # make export-marts
    python scripts/export_marts.py --check             # make export-check
    python scripts/export_marts.py --check --offline   # the files only, no Snowflake (works after the trial)

Why: the Snowflake trial ends about 2026-11-12. From then on the public dashboard serves a frozen
Google Sheets snapshot of the same MART tables (docs/architecture.md, decision 8; the steps are in
docs/sheets_snapshot.md). This script produces that snapshot from exactly what the dashboard reads:
it connects as SVC_CRM_BI with role CRM_REPORTER (key pair, settings from .env), the identity the
Data Studio connector uses, which can read MART and nothing else.

Export (default) into exports/. All or nothing: every file is written to a temporary directory
inside exports/ first, checked, and only then moved into place.
  <TABLE>.csv   one per table in EXPORT_TABLES. UTF-8, LF line ends, header = the exact UPPERCASE
                column names in table order, rows ORDER BY the grain key (checked unique). Numbers
                are plain decimals at the column's scale (NUMBER(18,6) -> 0.016258; never a float
                artefact or an exponent), dates YYYY-MM-DD, booleans TRUE/FALSE, NULL an empty field.
  marketing_crm_dashboard_snapshot.xlsx
                sheet README (provenance, export time, row counts, caveats, CC BY 4.0 attribution),
                then one sheet per table, named after it: numeric cells numeric (written as their
                exact decimal text, formatted at the column's scale), date cells dates
                (yyyy-mm-dd), header row bold and frozen, the attribution as a note on cell A1.
  manifest.json per table: rows, columns and types, SHA-256 of the CSV, the Snowflake query IDs,
                export time, source object, and the checksums COUNT(*), HASH_AGG(*), the SUM of
                every NUMBER column and the NULL count of every column; plus the dataset version and
                the dbt project.
The rows fetched are compared with a second query on the same table (COUNT, column sums, NULL
counts) before anything is written, so a table rebuilt in the middle of an export fails the run.

Check (--check), read-only:
  files      every CSV's SHA-256, header, row count, key order, column sums and NULL counts vs the
             manifest; every workbook sheet cell by cell vs its CSV (value, cell type, number
             format), header bold and frozen, README first, the attribution note on A1
  Snowflake  (skipped with --offline) as the BI user again: the columns and types, COUNT(*),
             HASH_AGG(*), column sums and NULL counts of each live table vs the manifest
Prints PASS / FAIL / WARN / INFO lines; exits 1 if anything differs, 2 if the configuration is
missing or Snowflake refuses the connection. Never prints key material, .env values or the account.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
import platform
import re
import shutil
import sys
import tempfile
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, Inexact, InvalidOperation, localcontext
from importlib import metadata
from pathlib import Path
from typing import Any, Iterator

sys.path.insert(0, str(Path(__file__).resolve().parent))
import sfconn  # noqa: E402  (stdlib-only at import time)

REPO_ROOT = sfconn.REPO_ROOT
DEFAULT_OUT_DIR = REPO_ROOT / "exports"
DBT_PROJECT_FILE = REPO_ROOT / "dbt" / "dbt_project.yml"
SCHEMA = "MART"
WORKBOOK_NAME = "marketing_crm_dashboard_snapshot.xlsx"
MANIFEST_NAME = "manifest.json"
MANIFEST_VERSION = 1
README_SHEET = "README"
TMP_PREFIX = ".export-"  # git-ignored (exports/.export-*/): only left behind by a crash


@dataclass(frozen=True)
class ExportTable:
    name: str
    order_by: tuple[str, ...]  # the grain key: unique, so the row order is deterministic
    grain: str
    basis: str


# Every MART table the dashboard reads, in the order of the workbook. To export another table,
# add it here (with its grain key), run `make export-marts`, then `make export-check`.
EXPORT_TABLES: tuple[ExportTable, ...] = (
    ExportTable("MART_CHANNEL_PERFORMANCE", ("CHANNEL",),
                "one row per channel",
                "full range 2023-01-01..2025-12-31, no date dimension; its ratio columns are valid "
                "at channel grain only"),
    ExportTable("MART_CAMPAIGN_PERFORMANCE", ("CAMPAIGN_ID",),
                "one row per campaign",
                "full range, no date dimension (START_DATE / START_MONTH are campaign attributes, "
                "not activity dates); its ratio columns are valid at campaign grain only"),
    ExportTable("MART_MONTHLY_MARKETING", ("MONTH_START", "CHANNEL"),
                "one row per activity month x channel (36 x 5)",
                "activity basis: MONTH_START; never on one date axis with COHORT_MONTH"),
    ExportTable("MART_MONTHLY_COHORT", ("COHORT_MONTH", "CHANNEL"),
                "one row per lead-cohort month x channel (36 x 5)",
                "lead-cohort basis: COHORT_MONTH; the only monthly source of conversion rates"),
    ExportTable("MART_CONVERSION_FUNNEL", ("CHANNEL", "STAGE_ORDER"),
                "one row per channel x funnel stage (5 x 5)",
                "full range; aggregate campaign funnel, STAGE_VALUE adds up across channels within "
                "a stage only"),
    ExportTable("MART_LEAD_QUALITY", ("CHANNEL", "SCORE_BAND"),
                "one row per channel x 10-point lead-score band (5 x 10)",
                "lead-cohort basis, full range, no date dimension; association only"),
    ExportTable("FCT_CAMPAIGN_DAILY_ACTIVITY", ("CAMPAIGN_ID", "ACTIVITY_DATE"),
                "one row per campaign x day, START_DATE to END_DATE + 21 days",
                "activity basis: ACTIVITY_DATE / ACTIVITY_MONTH; never divide CONVERSIONS by LEADS"),
    ExportTable("DIM_CAMPAIGN", ("CAMPAIGN_ID",),
                "one row per campaign",
                "campaign attributes (names, channel, objective, window, budget) for filters"),
)

DATASET = {
    "name": "Marketing Campaigns & Conversions Dataset",
    "version": "v1.0.0",
    "publisher": "Analytics Engineering",
    "url": "https://www.analyticsengineering.com/datasets/marketing-campaigns",
    "license": "CC BY 4.0",
    "license_url": "https://creativecommons.org/licenses/by/4.0/",
    "synthetic": True,
}
ATTRIBUTION = (
    'Data: Analytics Engineering, "Marketing Campaigns & Conversions Dataset v1.0.0" '
    "(https://www.analyticsengineering.com/datasets/marketing-campaigns), CC BY 4.0 "
    "(https://creativecommons.org/licenses/by/4.0/). Typed, cleaned and aggregated by Hoedong Kim; "
    "not endorsed by the publisher."
)
CAVEATS = (
    "Synthetic data: a fictional business, no real companies or people, no industry benchmarks.",
    "Revenue is last-touch attributed: ROAS and revenue by channel or campaign describe "
    "attribution, not the causal effect of spend. Offline and view-through conversions are not included.",
    "Two time bases, never on one date axis: activity months (MART_MONTHLY_MARKETING.MONTH_START, "
    "FCT_CAMPAIGN_DAILY_ACTIVITY.ACTIVITY_DATE) and lead-cohort months "
    "(MART_MONTHLY_COHORT.COHORT_MONTH). Conversion rates exist on the cohort basis only.",
    "A ratio over several rows is SUM(numerator) / SUM(denominator); never SUM or AVG a stored "
    "ratio column. Budget utilisation only over whole campaigns (campaign or channel tables).",
    "Every KPI is computed in Snowflake SQL (dbt models); this file only carries the results. "
    "Definitions: docs/kpi_definitions.md and docs/data_dictionary.md in the repository.",
)
NOTE_AUTHOR = "Hoedong Kim"

IDENTIFIER = re.compile(r"[A-Z][A-Z0-9_]*")
SUPPORTED_TYPES = {"NUMBER", "TEXT", "DATE", "BOOLEAN"}
MAX_SIGNIFICANT_DIGITS = 15  # what a spreadsheet cell (an IEEE double) holds exactly


class ExportError(Exception):
    """The export cannot be written correctly. The message is safe to print."""


# ---------------------------------------------------------------------------------------------
# Columns and value formatting
# ---------------------------------------------------------------------------------------------

@dataclass(frozen=True)
class Column:
    name: str
    data_type: str  # NUMBER, TEXT, DATE or BOOLEAN (INFORMATION_SCHEMA.COLUMNS.DATA_TYPE)
    precision: int | None = None
    scale: int | None = None

    @property
    def label(self) -> str:
        if self.data_type == "NUMBER":
            return f"NUMBER({self.precision},{self.scale})"
        return self.data_type

    @property
    def is_number(self) -> bool:
        return self.data_type == "NUMBER"

    @classmethod
    def from_label(cls, name: str, label: str) -> Column:
        match = re.fullmatch(r"NUMBER\((\d+),(\d+)\)", label)
        if match:
            return cls(name, "NUMBER", int(match.group(1)), int(match.group(2)))
        if label not in SUPPORTED_TYPES:
            raise ExportError(f"column {name}: unsupported type {label} in the manifest")
        return cls(name, label)


def quantize(value: Decimal, scale: int) -> Decimal:
    """value at exactly `scale` decimal places; raises ExportError if that would round."""
    with localcontext() as ctx:
        ctx.prec = 80
        ctx.traps[Inexact] = True
        try:
            result = value.quantize(Decimal(1).scaleb(-scale))
        except (Inexact, InvalidOperation):
            raise ExportError(f"{value} has more than {scale} decimal places") from None
    return result.copy_abs() if result.is_zero() else result  # never "-0.00"


def to_decimal(value: Any) -> Decimal:
    if isinstance(value, bool) or not isinstance(value, (int, Decimal)):
        raise ExportError(f"expected an exact number, got {type(value).__name__} {value!r}")
    return Decimal(value)


def format_number(value: Any, scale: int) -> str:
    """Plain decimal text at the column's scale: 24, 857835.07, 0.016258, 0.000000."""
    return format(quantize(to_decimal(value), scale), "f")


def render(value: Any, col: Column) -> str:
    """The CSV text of one value."""
    if value is None:
        return ""
    if col.data_type == "NUMBER":
        return format_number(value, col.scale)
    if col.data_type == "DATE":
        if isinstance(value, datetime) or not isinstance(value, date):
            raise ExportError(f"column {col.name}: expected a date, got {type(value).__name__}")
        return value.isoformat()
    if col.data_type == "BOOLEAN":
        if not isinstance(value, bool):
            raise ExportError(f"column {col.name}: expected a boolean, got {type(value).__name__}")
        return "TRUE" if value else "FALSE"
    if not isinstance(value, str):
        raise ExportError(f"column {col.name}: expected text, got {type(value).__name__}")
    return value


def significant_digits(value: Decimal) -> int:
    digits = value.normalize().as_tuple().digits if not value.is_zero() else (0,)
    return len(digits)


def number_format(col: Column) -> str | None:
    """Spreadsheet number format: the column's scale, no thousands separator (nothing to parse)."""
    if col.data_type == "NUMBER":
        return "0" if col.scale == 0 else "0." + "0" * col.scale
    if col.data_type == "DATE":
        return "yyyy-mm-dd"
    return None


# ---------------------------------------------------------------------------------------------
# Statistics shared by the export and the check
# ---------------------------------------------------------------------------------------------

class Stats:
    """Row count, exact SUM of every NUMBER column and NULL count of every column."""

    def __init__(self, columns: list[Column]) -> None:
        self.columns = columns
        self.rows = 0
        self.sums = {c.name: Decimal(0) for c in columns if c.is_number}
        self.nulls = {c.name: 0 for c in columns}

    def add(self, values: list[Any]) -> None:
        """values: Python values (export) or Decimal / str / None parsed from the CSV (check)."""
        self.rows += 1
        for col, value in zip(self.columns, values):
            if value is None:
                self.nulls[col.name] += 1
            elif col.is_number:
                self.sums[col.name] += to_decimal(value)

    def as_manifest(self) -> dict[str, Any]:
        by_name = {c.name: c for c in self.columns}
        return {
            "rows": self.rows,
            "sums": {n: format_number(s, by_name[n].scale) for n, s in self.sums.items()},
            "null_counts": dict(self.nulls),
        }


def aggregate_sql(fqn: str, columns: list[Column]) -> str:
    parts = ["COUNT(*)", "HASH_AGG(*)"]
    parts += [f'SUM("{c.name}")' for c in columns if c.is_number]
    parts += [f'COUNT_IF("{c.name}" IS NULL)' for c in columns]
    return f"SELECT {', '.join(parts)} FROM {fqn}"


def parse_aggregate(row: tuple, columns: list[Column]) -> dict[str, Any]:
    """The aggregate_sql row as manifest checksums (an all-NULL column sums to 0)."""
    values = list(row)
    count, hash_agg = values[0], values[1]
    numbers = [c for c in columns if c.is_number]
    sums_raw = values[2:2 + len(numbers)]
    nulls_raw = values[2 + len(numbers):]
    return {
        "rows": int(count),
        "hash_agg": str(hash_agg),
        "sums": {c.name: format_number(0 if v is None else v, c.scale) for c, v in zip(numbers, sums_raw)},
        "null_counts": {c.name: int(v) for c, v in zip(columns, nulls_raw)},
    }


# ---------------------------------------------------------------------------------------------
# Snowflake side
# ---------------------------------------------------------------------------------------------

def connect_bi(tool: str):
    """Connect as the BI user with the reporter role; returns (connection, settings)."""
    sfconn.load_env()
    settings = sfconn.resolve_settings(
        role=sfconn.setting("SNOWFLAKE_REPORTER_ROLE"),
        user_var="SNOWFLAKE_BI_USER",
        key_var="SNOWFLAKE_BI_PRIVATE_KEY_PATH",
        passphrase_var=None,
        schema=SCHEMA,
        allow_encrypted=False,
    )
    conn = sfconn.connect(settings, tool=tool, numbers_as_decimal=True)
    cur = conn.cursor()
    try:
        cur.execute("USE SECONDARY ROLES NONE")
        cur.execute("SELECT CURRENT_USER(), CURRENT_ROLE(), CURRENT_WAREHOUSE(), CURRENT_DATABASE()")
        user, role, warehouse, database = cur.fetchone()
    finally:
        cur.close()
    if str(user).upper() != settings.user.upper() or str(role).upper() != settings.role.upper():
        conn.close()
        raise ExportError(f"session is {user} / {role}, expected {settings.user} / {settings.role}")
    print(f"INFO  connected as {user} / {role} (warehouse {warehouse}, database {database}), "
          "the Data Studio connector's identity")
    return conn, settings


def database_name(settings) -> str:
    name = settings.database.upper()
    if not IDENTIFIER.fullmatch(name):
        raise ExportError(f"SNOWFLAKE_DATABASE {settings.database!r} is not a plain identifier")
    return name


def read_metadata(conn, database: str, names: list[str]) -> tuple[dict, dict]:
    """INFORMATION_SCHEMA rows of the tables, as CRM_REPORTER sees them."""
    bad = [n for n in names if not IDENTIFIER.fullmatch(n)]
    if bad:
        raise ExportError(f"not plain UPPERCASE table names: {bad}")
    in_list = ", ".join(f"'{n}'" for n in names)
    cur = conn.cursor()
    try:
        cur.execute(
            f"SELECT TABLE_NAME, LAST_ALTERED, COMMENT "
            f"FROM {database}.INFORMATION_SCHEMA.TABLES "
            f"WHERE TABLE_SCHEMA = '{SCHEMA}' AND TABLE_NAME IN ({in_list})"
        )
        tables = {r[0]: {"last_altered": r[1], "comment": r[2]} for r in cur.fetchall()}
        cur.execute(
            f"SELECT TABLE_NAME, COLUMN_NAME, DATA_TYPE, NUMERIC_PRECISION, NUMERIC_SCALE "
            f"FROM {database}.INFORMATION_SCHEMA.COLUMNS "
            f"WHERE TABLE_SCHEMA = '{SCHEMA}' AND TABLE_NAME IN ({in_list}) "
            f"ORDER BY TABLE_NAME, ORDINAL_POSITION"
        )
        columns: dict[str, list[Column]] = {}
        for table, name, data_type, precision, scale in cur.fetchall():
            col = Column(name, data_type,
                         int(precision) if data_type == "NUMBER" else None,
                         int(scale) if data_type == "NUMBER" else None)
            columns.setdefault(table, []).append(col)
    finally:
        cur.close()
    return tables, columns


def validate_columns(table: ExportTable, columns: list[Column]) -> None:
    names = [c.name for c in columns]
    for col in columns:
        if not IDENTIFIER.fullmatch(col.name):
            raise ExportError(f"{table.name}.{col.name}: not an UPPERCASE ASCII name (Data Studio field names)")
        if col.data_type not in SUPPORTED_TYPES:
            raise ExportError(f"{table.name}.{col.name}: type {col.label} is not exported "
                              f"(supported: {', '.join(sorted(SUPPORTED_TYPES))})")
    missing = [k for k in table.order_by if k not in names]
    if missing:
        raise ExportError(f"{table.name}: grain key column(s) {missing} not in the table")


def utc_text(value: datetime | None) -> str | None:
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def now_utc() -> str:
    return utc_text(datetime.now(timezone.utc))


# ---------------------------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------------------------

@dataclass
class ExportedTable:
    spec: ExportTable
    fqn: str
    columns: list[Column]
    rows: list[tuple]
    csv_bytes: bytes
    query_id: str
    checksum_query_id: str
    exported_at: str
    checksums: dict[str, Any]
    last_altered: str | None
    comment: str | None


def fetch_table(conn, database: str, spec: ExportTable, columns: list[Column], meta: dict) -> ExportedTable:
    quoted = f'"{database}"."{SCHEMA}"."{spec.name}"'
    select_list = ", ".join(f'"{c.name}"' for c in columns)
    order_list = ", ".join(f'"{k}"' for k in spec.order_by)
    cur = conn.cursor()
    try:
        exported_at = now_utc()
        cur.execute(f"SELECT {select_list} FROM {quoted} ORDER BY {order_list}")
        query_id = cur.sfqid
        rows = cur.fetchall()
        cur.execute(aggregate_sql(quoted, columns))
        checksum_query_id = cur.sfqid
        sql_checks = parse_aggregate(cur.fetchone(), columns)
    finally:
        cur.close()

    # Render the CSV, collect the same statistics locally, and check the grain key.
    stats = Stats(columns)
    key_positions = [[c.name for c in columns].index(k) for k in spec.order_by]
    seen_keys: set[tuple] = set()
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow([c.name for c in columns])
    for row in rows:
        try:
            writer.writerow([render(v, c) for v, c in zip(row, columns)])
            for v, c in zip(row, columns):
                if c.is_number and v is not None and significant_digits(to_decimal(v)) > MAX_SIGNIFICANT_DIGITS:
                    raise ExportError(f"column {c.name}: {v} has more than {MAX_SIGNIFICANT_DIGITS} "
                                      "significant digits, more than a spreadsheet cell holds")
        except ExportError as exc:
            raise ExportError(f"{spec.name}: {exc}") from None
        stats.add(list(row))
        key = tuple(row[i] for i in key_positions)
        if any(k is None for k in key):
            raise ExportError(f"{spec.name}: NULL in the grain key {spec.order_by}")
        if key in seen_keys:
            raise ExportError(f"{spec.name}: grain key {spec.order_by} is not unique ({key})")
        seen_keys.add(key)

    local = stats.as_manifest()
    for field in ("rows", "sums", "null_counts"):
        if local[field] != sql_checks[field]:
            raise ExportError(
                f"{spec.name}: the {field} of the rows fetched differ from a second query on the "
                "table; was it rebuilt during the export? Run the export again."
            )
    return ExportedTable(
        spec=spec, fqn=f"{database}.{SCHEMA}.{spec.name}", columns=columns, rows=rows,
        csv_bytes=buffer.getvalue().encode("utf-8"), query_id=query_id,
        checksum_query_id=checksum_query_id, exported_at=exported_at, checksums=sql_checks,
        last_altered=utc_text(meta.get("last_altered")), comment=meta.get("comment"),
    )


@contextmanager
def exact_decimal_cells() -> Iterator[None]:
    """openpyxl writes every number as '%.16g' % value, which turns Decimal('857835.07') into
    857835.0699999999 in the sheet XML. While saving, write Decimal and int cells as their exact
    decimal text instead (both of openpyxl's cell writers look safe_string up at call time)."""
    import openpyxl.cell._writer as cell_writer

    original = cell_writer.safe_string

    def safe_string(value: Any) -> str:
        if isinstance(value, Decimal):
            return format(value, "f")
        if type(value) is int:
            return str(value)
        return original(value)

    cell_writer.safe_string = safe_string
    try:
        yield
    finally:
        cell_writer.safe_string = original


def cell_value(value: Any, col: Column) -> Any:
    if value is None:
        return None
    if col.is_number:
        return quantize(to_decimal(value), col.scale) if col.scale else int(value)
    return value


def column_width(header: str, texts: list[str]) -> float:
    longest = max([len(header)] + [len(t) for t in texts[:500]])
    return min(max(longest + 2, 8), 48)


def sheet_note(table: ExportedTable, exported_at: str) -> str:
    return (f"{table.fqn}, exported {exported_at} as SVC_CRM_BI / CRM_REPORTER; "
            f"{len(table.rows)} rows. {ATTRIBUTION}")


def write_workbook(path: Path, tables: list[ExportedTable], context: dict[str, Any]) -> None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font

    bold = Font(bold=True)
    wb = Workbook()
    readme = wb.active
    readme.title = README_SHEET
    write_readme_sheet(readme, tables, context, bold)

    for table in tables:
        if len(table.spec.name) > 31:
            raise ExportError(f"{table.spec.name}: longer than 31 characters, the sheet name limit")
        ws = wb.create_sheet(title=table.spec.name)
        ws.append([c.name for c in table.columns])
        for cell in ws[1]:
            cell.font = bold
        for row in table.rows:
            ws.append([cell_value(v, c) for v, c in zip(row, table.columns)])
        csv_rows = list(csv.reader(io.StringIO(table.csv_bytes.decode("utf-8"), newline="")))[1:]
        for index, col in enumerate(table.columns, start=1):
            fmt = number_format(col)
            column_cells = ws.iter_rows(min_row=2, max_row=len(table.rows) + 1,
                                        min_col=index, max_col=index)
            for (cell,) in column_cells:
                if fmt:
                    cell.number_format = fmt
                if col.data_type == "TEXT" and cell.value is not None:
                    cell.data_type = "s"  # never a formula, whatever the text starts with
            letter = ws.cell(row=1, column=index).column_letter
            ws.column_dimensions[letter].width = column_width(col.name, [r[index - 1] for r in csv_rows])
        ws.freeze_panes = "A2"
        note = Comment(sheet_note(table, context["exported_at_utc"]), NOTE_AUTHOR)
        note.width, note.height = 420, 160
        ws["A1"].comment = note
        ws["A1"].alignment = Alignment(vertical="top")

    props = wb.properties
    props.title = "Marketing CRM & Conversion Funnel Analytics: dashboard data snapshot"
    props.creator = NOTE_AUTHOR
    props.description = ATTRIBUTION
    created = datetime.fromisoformat(context["exported_at_utc"].replace("Z", "+00:00")).replace(tzinfo=None)
    props.created = props.modified = created
    with exact_decimal_cells():
        wb.save(path)


def write_readme_sheet(ws, tables: list[ExportedTable], context: dict[str, Any], bold) -> None:
    from openpyxl.styles import Alignment, Font

    ws.append(["Marketing CRM & Conversion Funnel Analytics: dashboard data snapshot"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    rows = [
        ("What this is", "A frozen export of the MART tables the Google Data Studio dashboard reads. "
                         "Every sheet after this one is one table, named after it; row 1 holds its "
                         "exact UPPERCASE column names and the data starts in row 2. Do not edit, "
                         "sort or add cells to those sheets: each one is a Data Studio data source."),
        ("Source", f"Snowflake {context['database']}.{SCHEMA}: tables built by dbt Core "
                   f"(project {context['dbt_project']['name']} {context['dbt_project']['version']}), "
                   "read as SVC_CRM_BI / CRM_REPORTER, the dashboard connector's read-only identity."),
        ("Exported at (UTC)", context["exported_at_utc"]),
        ("Produced by", "scripts/export_marts.py (make export-marts). exports/manifest.json holds each "
                        "table's row count, SHA-256 of its CSV, Snowflake query IDs, HASH_AGG and "
                        "column sums; make export-check re-verifies them."),
        ("Dataset", f"{DATASET['publisher']}, \"{DATASET['name']} {DATASET['version']}\", "
                    f"{DATASET['url']} (synthetic)"),
        ("Attribution", ATTRIBUTION),
        ("License", f"Data: {DATASET['license']} ({DATASET['license_url']}). "
                    "Code in the repository: MIT."),
    ]
    rows += [(f"Caveat {index}", caveat) for index, caveat in enumerate(CAVEATS, start=1)]
    for row_index, (label, text) in enumerate(rows, start=3):
        ws.cell(row=row_index, column=1, value=label).font = bold
        ws.cell(row=row_index, column=2, value=text)
    header_row = 3 + len(rows) + 1  # one empty row after the notes
    headers = ["Sheet", "Rows", "Columns", "Grain", "Basis and use", "Snowflake object",
               "Last altered in Snowflake (UTC)", "CSV SHA-256"]
    for col_index, text in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col_index, value=text).font = bold
    for row_index, table in enumerate(tables, start=header_row + 1):
        values = [table.spec.name, len(table.rows), len(table.columns), table.spec.grain,
                  table.spec.basis, table.fqn, table.last_altered,
                  hashlib.sha256(table.csv_bytes).hexdigest()]
        for col_index, value in enumerate(values, start=1):
            ws.cell(row=row_index, column=col_index, value=value)
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for letter, width in zip("ABCDEFGH", (30, 110, 9, 34, 48, 44, 22, 66)):
        ws.column_dimensions[letter].width = width


def build_manifest(tables: list[ExportedTable], context: dict[str, Any], workbook_sha: str) -> dict:
    return {
        "manifest_version": MANIFEST_VERSION,
        "description": "Frozen export of the MART tables the Data Studio dashboard reads "
                       "(Google Sheets snapshot, docs/sheets_snapshot.md).",
        "exported_at_utc": context["exported_at_utc"],
        "source": {
            "platform": "Snowflake",
            "database": context["database"],
            "schema": SCHEMA,
            "built_by": "dbt Core",
            "read_as_user": context["user"],
            "read_as_role": context["role"],
        },
        "dbt_project": context["dbt_project"],
        "dataset": DATASET,
        "attribution": ATTRIBUTION,
        "caveats": list(CAVEATS),
        "tools": context["tools"],
        "csv_format": "UTF-8, LF, header = exact column names; numbers as plain decimals at the "
                      "column scale; dates YYYY-MM-DD; booleans TRUE/FALSE; NULL = empty field",
        "workbook": {
            "file": WORKBOOK_NAME,
            "sha256": workbook_sha,
            "sheets": [README_SHEET] + [t.spec.name for t in tables],
        },
        "tables": [
            {
                "table": t.spec.name,
                "source_object": t.fqn,
                "csv": f"{t.spec.name}.csv",
                "sheet": t.spec.name,
                "rows": len(t.rows),
                "columns": [{"name": c.name, "type": c.label} for c in t.columns],
                "order_by": list(t.spec.order_by),
                "grain": t.spec.grain,
                "basis": t.spec.basis,
                "sha256": hashlib.sha256(t.csv_bytes).hexdigest(),
                "csv_bytes": len(t.csv_bytes),
                "snowflake_query_id": t.query_id,
                "checksum_query_id": t.checksum_query_id,
                "exported_at_utc": t.exported_at,
                "table_last_altered_utc": t.last_altered,
                "table_comment": t.comment,
                "checksums": t.checksums,
            }
            for t in tables
        ],
    }


def dbt_project_info() -> dict[str, str | None]:
    text = DBT_PROJECT_FILE.read_text(encoding="utf-8") if DBT_PROJECT_FILE.is_file() else ""
    name = re.search(r"^name:\s*[\"']?([A-Za-z0-9_]+)", text, flags=re.MULTILINE)
    version = re.search(r"^version:\s*[\"']?([0-9A-Za-z._-]+)", text, flags=re.MULTILINE)
    return {"name": name.group(1) if name else None,
            "version": version.group(1) if version else None,
            "path": "dbt/dbt_project.yml"}


def tool_versions() -> dict[str, str]:
    versions = {"python": platform.python_version()}
    for package in ("snowflake-connector-python", "openpyxl", "dbt-core"):
        try:
            versions[package] = metadata.version(package)
        except metadata.PackageNotFoundError:
            pass
    return versions


def rel(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(REPO_ROOT))
    except ValueError:
        return str(path)


def run_export(out_dir: Path) -> int:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("FAIL  openpyxl is not installed. Run 'make setup' first.")
        return 2
    try:
        conn, settings = connect_bi("export_marts")
    except (sfconn.ConfigError, sfconn.ConnectError, ExportError) as exc:
        print(f"FAIL  {exc}")
        return 2
    from snowflake.connector.errors import Error as SnowflakeError

    names = [t.name for t in EXPORT_TABLES]
    try:
        database = database_name(settings)
        meta, columns_by_table = read_metadata(conn, database, names)
        missing = [n for n in names if n not in meta or n not in columns_by_table]
        if missing:
            raise ExportError(f"not visible to {settings.role} in {database}.{SCHEMA}: {missing} "
                              "(build them with dbt first)")
        context = {
            "exported_at_utc": now_utc(),
            "database": database,
            "user": settings.user,
            "role": settings.role,
            "dbt_project": dbt_project_info(),
            "tools": tool_versions(),
        }
        exported = []
        for spec in EXPORT_TABLES:
            columns = columns_by_table[spec.name]
            validate_columns(spec, columns)
            table = fetch_table(conn, database, spec, columns, meta[spec.name])
            exported.append(table)
            print(f"PASS  {spec.name}: {len(table.rows)} rows x {len(columns)} columns fetched; row "
                  f"count, {len(table.checksums['sums'])} column sums and NULL counts equal a second "
                  f"query (data {table.query_id}, checksums {table.checksum_query_id})")
    except ExportError as exc:
        print(f"FAIL  {exc}")
        return 1
    except SnowflakeError as exc:  # a failed query: the connector's message, never credentials
        print(f"FAIL  {sfconn.describe_error(exc)}")
        return 1
    finally:
        conn.close()

    out_dir.mkdir(parents=True, exist_ok=True)
    tmp = Path(tempfile.mkdtemp(prefix=TMP_PREFIX, dir=out_dir))
    try:
        for table in exported:
            (tmp / f"{table.spec.name}.csv").write_bytes(table.csv_bytes)
        workbook = tmp / WORKBOOK_NAME
        try:
            write_workbook(workbook, exported, context)
        except ExportError as exc:
            print(f"FAIL  {exc}")
            return 1
        manifest = build_manifest(exported, context, hashlib.sha256(workbook.read_bytes()).hexdigest())
        (tmp / MANIFEST_NAME).write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")

        print("INFO  verifying the new files before they replace the old ones")
        report = Report()
        check_files(report, tmp, verbose=False)
        if report.failures:
            print(f"FAIL  the new files failed {report.failures} check(s); {rel(out_dir)} left unchanged")
            return 1
        produced = [f"{t.spec.name}.csv" for t in exported] + [WORKBOOK_NAME, MANIFEST_NAME]
        for name in produced:
            os.replace(tmp / name, out_dir / name)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    stale = sorted(p.name for p in out_dir.glob("*.csv") if p.name not in set(produced))
    if stale:
        print(f"WARN  {rel(out_dir)} also holds CSVs that are not in this export: {stale}")
    total_rows = sum(len(t.rows) for t in exported)
    print(f"PASS  wrote {len(exported)} CSVs ({total_rows} rows), {WORKBOOK_NAME} "
          f"({len(exported) + 1} sheets) and {MANIFEST_NAME} to {rel(out_dir)}/ "
          f"(exported {context['exported_at_utc']})")
    return 0


# ---------------------------------------------------------------------------------------------
# Check
# ---------------------------------------------------------------------------------------------

class Report:
    def __init__(self) -> None:
        self.failures = 0
        self.warnings = 0

    def line(self, status: str, name: str, detail: str = "", *, show: bool = True) -> None:
        if show:
            print(f"{status:<5} {name}" + (f": {detail}" if detail else ""))

    def check(self, ok: bool, name: str, ok_detail: str, fail_detail: str, *, show: bool = True) -> bool:
        if not ok:
            self.failures += 1
            self.line("FAIL", name, fail_detail)  # failures always print
        else:
            self.line("PASS", name, ok_detail, show=show)
        return ok

    def warn(self, name: str, detail: str) -> None:
        self.warnings += 1
        self.line("WARN", name, detail)


def load_manifest(out_dir: Path) -> dict:
    path = out_dir / MANIFEST_NAME
    if not path.is_file():
        raise ExportError(f"{rel(path)} does not exist: run 'make export-marts' first")
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ExportError(f"{rel(path)} is not valid JSON: {exc}") from None


def manifest_columns(entry: dict) -> list[Column]:
    columns = [Column.from_label(c["name"], c["type"]) for c in entry["columns"]]
    bad = [c.name for c in columns if not IDENTIFIER.fullmatch(c.name)]
    if bad:
        raise ExportError(f"{entry.get('table')}: not plain UPPERCASE column names in the manifest: {bad}")
    return columns


def parse_csv_value(text: str, col: Column) -> Any:
    """CSV text -> Decimal for numbers (must be canonical), the text otherwise, None if empty."""
    if text == "":
        return None
    if col.is_number:
        try:
            value = Decimal(text)
        except InvalidOperation:
            raise ExportError(f"column {col.name}: {text!r} is not a number") from None
        if format_number(value, col.scale) != text:
            raise ExportError(f"column {col.name}: {text!r} is not a plain decimal at scale {col.scale}")
        return value
    if col.data_type == "DATE" and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        raise ExportError(f"column {col.name}: {text!r} is not a YYYY-MM-DD date")
    if col.data_type == "BOOLEAN" and text not in ("TRUE", "FALSE"):
        raise ExportError(f"column {col.name}: {text!r} is not TRUE or FALSE")
    return text


def check_csv(report: Report, out_dir: Path, entry: dict, verbose: bool) -> list[list[str]] | None:
    name = entry["table"]
    path = out_dir / entry["csv"]
    if not report.check(path.is_file(), f"{name} csv", "", f"{rel(path)} is missing", show=False):
        return None
    failures_before = report.failures
    data = path.read_bytes()
    sha = hashlib.sha256(data).hexdigest()
    report.check(sha == entry["sha256"], f"{name} csv sha256", "",
                 f"{sha} != manifest {entry['sha256']}", show=False)
    columns = manifest_columns(entry)
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        report.check(False, f"{name} csv", "", "not UTF-8")
        return None
    rows = list(csv.reader(io.StringIO(text, newline="")))
    if not report.check(bool(rows) and rows[0] == [c.name for c in columns], f"{name} csv header",
                        "", f"{rows[0] if rows else 'empty file'} != manifest columns", show=False):
        return None
    body = rows[1:]
    stats = Stats(columns)
    key_index = [[c.name for c in columns].index(k) for k in entry["order_by"]]
    previous_key = None
    problems: list[str] = []
    for line_no, row in enumerate(body, start=2):
        if len(row) != len(columns):
            problems.append(f"line {line_no}: {len(row)} fields")
            continue
        try:
            values = [parse_csv_value(t, c) for t, c in zip(row, columns)]
        except ExportError as exc:
            problems.append(f"line {line_no}: {exc}")
            continue
        stats.add(values)
        key = tuple(values[i] for i in key_index)
        if any(k is None for k in key):
            problems.append(f"line {line_no}: NULL in the grain key")
            continue
        if previous_key is not None and not previous_key < key:
            problems.append(f"line {line_no}: key {key} not after {previous_key}")
        previous_key = key
    report.check(not problems, f"{name} csv values", "",
                 f"{len(problems)} problem(s), first: {problems[:3]}", show=False)
    local = stats.as_manifest()
    expected = entry["checksums"]
    report.check(local["rows"] == entry["rows"] == expected["rows"], f"{name} csv rows", "",
                 f"{local['rows']} in the CSV, manifest rows {entry['rows']} / "
                 f"checksum rows {expected['rows']}", show=False)
    sum_diffs = {k: (v, expected["sums"].get(k)) for k, v in local["sums"].items() if v != expected["sums"].get(k)}
    report.check(not sum_diffs and set(local["sums"]) == set(expected["sums"]), f"{name} csv sums", "",
                 f"differ (csv, manifest): {sum_diffs or 'column set differs'}", show=False)
    report.check(local["null_counts"] == expected["null_counts"], f"{name} csv NULLs", "",
                 f"{local['null_counts']} != {expected['null_counts']}", show=False)
    if report.failures == failures_before:
        report.line("PASS", f"{name} csv",
                    f"SHA-256 {sha[:12]}..., {len(columns)} columns, {local['rows']} rows strictly ordered "
                    f"by {', '.join(entry['order_by'])}, every field at its column type, "
                    f"{len(local['sums'])} column sums and NULL counts = manifest", show=verbose)
    return body


def xlsx_text(cell, col: Column) -> tuple[str, str | None]:
    """(the CSV text a workbook cell stands for, a problem or None)."""
    value = cell.value
    if value is None:
        return "", None
    if col.is_number:
        if isinstance(value, bool) or not isinstance(value, (int, float)) or cell.data_type != "n":
            return repr(value), "not a numeric cell"
        if cell.number_format != number_format(col):
            return repr(value), f"number format {cell.number_format!r}"
        try:
            exact = Decimal(repr(value)) if isinstance(value, float) else Decimal(value)
            return format_number(exact, col.scale), None
        except ExportError:
            return repr(value), "not exact at the column scale"
    if col.data_type == "DATE":
        if not isinstance(value, datetime) or not cell.is_date:
            return repr(value), "not a date cell"
        if cell.number_format != number_format(col):
            return repr(value), f"number format {cell.number_format!r}"
        return value.date().isoformat(), None
    if col.data_type == "BOOLEAN":
        if not isinstance(value, bool):
            return repr(value), "not a boolean cell"
        return ("TRUE" if value else "FALSE"), None
    if not isinstance(value, str) or cell.data_type != "s":
        return repr(value), "not a text cell"
    return value, None


def check_workbook(report: Report, out_dir: Path, manifest: dict, csv_bodies: dict, verbose: bool) -> None:
    from openpyxl import load_workbook

    path = out_dir / manifest["workbook"]["file"]
    if not report.check(path.is_file(), "workbook", "", f"{rel(path)} is missing", show=False):
        return
    sha = hashlib.sha256(path.read_bytes()).hexdigest()
    report.check(sha == manifest["workbook"]["sha256"], "workbook sha256", f"{sha[:12]}... = manifest",
                 f"{sha} != manifest {manifest['workbook']['sha256']}", show=verbose)
    wb = load_workbook(path)
    expected_sheets = [README_SHEET] + [t["sheet"] for t in manifest["tables"]]
    report.check(wb.sheetnames == expected_sheets == manifest["workbook"]["sheets"], "workbook sheets",
                 f"{README_SHEET} first, then {len(expected_sheets) - 1} table sheets",
                 f"{wb.sheetnames} != {expected_sheets}", show=verbose)
    readme_text = " ".join(str(c.value) for row in wb[README_SHEET].iter_rows() for c in row if c.value) \
        if README_SHEET in wb.sheetnames else ""
    report.check(ATTRIBUTION in readme_text and manifest["exported_at_utc"] in readme_text,
                 "workbook README", "attribution and export time present",
                 "the README sheet lacks the attribution or the export time", show=verbose)
    for entry in manifest["tables"]:
        name = entry["sheet"]
        if name not in wb.sheetnames or csv_bodies.get(entry["table"]) is None:
            report.check(False, f"{name} sheet", "", "missing sheet or unreadable CSV")
            continue
        ws = wb[name]
        columns = manifest_columns(entry)
        header = [c.value for c in ws[1]]
        layout_ok = (header == [c.name for c in columns] and all(c.font.b for c in ws[1])
                     and ws.freeze_panes == "A2" and ws.max_row == entry["rows"] + 1
                     and ws.max_column == len(columns))
        note = ws["A1"].comment.text if ws["A1"].comment else ""
        layout_ok = report.check(
            layout_ok and ATTRIBUTION in note, f"{name} sheet layout", "",
            f"header {header[:4]}..., bold {all(c.font.b for c in ws[1])}, freeze "
            f"{ws.freeze_panes}, size {ws.max_row}x{ws.max_column}, note {'yes' if note else 'no'}",
            show=False)
        problems: list[str] = []
        body = csv_bodies[entry["table"]]
        for cells, csv_row in zip(ws.iter_rows(min_row=2), body):
            for cell, col, expected in zip(cells, columns, csv_row):
                text, problem = xlsx_text(cell, col)
                if problem or text != expected:
                    problems.append(f"{cell.coordinate} {col.name}: {problem or f'{text!r} != csv {expected!r}'}")
        cells_ok = report.check(not problems, f"{name} sheet cells", "",
                                f"{len(problems)} cell(s) differ, first: {problems[:3]}", show=False)
        if layout_ok and cells_ok:
            report.line("PASS", f"{name} sheet",
                        f"{entry['rows'] * len(columns)} cells = the CSV (numeric / date / text cells, "
                        "number formats at the column scale); header bold and frozen; no stray cells; "
                        "attribution note on A1", show=verbose)


def check_files(report: Report, out_dir: Path, verbose: bool = True) -> dict | None:
    try:
        manifest = load_manifest(out_dir)
    except ExportError as exc:
        report.check(False, "manifest", "", str(exc))
        return None
    names = [t["table"] for t in manifest.get("tables", [])]
    report.check(names == [t.name for t in EXPORT_TABLES], "manifest tables",
                 f"{len(names)} tables, the EXPORT_TABLES list",
                 f"{names} != EXPORT_TABLES {[t.name for t in EXPORT_TABLES]}: run 'make export-marts'",
                 show=verbose)
    report.check(manifest.get("dataset", {}).get("version") == DATASET["version"]
                 and bool(manifest.get("dbt_project", {}).get("name")), "manifest provenance",
                 f"dataset {DATASET['version']}, dbt project {manifest.get('dbt_project', {}).get('name')}",
                 "dataset version or dbt project name missing", show=verbose)
    csv_bodies = {entry["table"]: check_csv(report, out_dir, entry, verbose) for entry in manifest["tables"]}
    check_workbook(report, out_dir, manifest, csv_bodies, verbose)
    return manifest


def check_live(report: Report, manifest: dict) -> None:
    conn, settings = connect_bi("export_check")
    try:
        database = database_name(settings)
        names = [t["table"] for t in manifest["tables"]]
        meta, columns_by_table = read_metadata(conn, database, names)
        cur = conn.cursor()
        try:
            for entry in manifest["tables"]:
                name = entry["table"]
                live_columns = columns_by_table.get(name)
                if not report.check(live_columns is not None, f"{name} live", "",
                                    f"not visible to {settings.role} in {database}.{SCHEMA}", show=False):
                    continue
                live_labels = [{"name": c.name, "type": c.label} for c in live_columns]
                report.check(live_labels == entry["columns"], f"{name} live columns",
                             f"{len(live_labels)} columns, names and types as exported",
                             "the table's columns or types changed since the export: "
                             f"{[c['name'] + ' ' + c['type'] for c in live_labels]}")
                cur.execute(aggregate_sql(f'"{database}"."{SCHEMA}"."{name}"', manifest_columns(entry)))
                live = parse_aggregate(cur.fetchone(), manifest_columns(entry))
                expected = entry["checksums"]
                diffs = [k for k in ("rows", "hash_agg", "sums", "null_counts") if live[k] != expected[k]]
                report.check(not diffs, f"{name} live checksums",
                             f"{live['rows']} rows, HASH_AGG {live['hash_agg']}, "
                             f"{len(live['sums'])} column sums and NULL counts = manifest = CSV "
                             f"(query {cur.sfqid})",
                             f"{', '.join(diffs)} differ from the manifest (live rows {live['rows']}, "
                             f"HASH_AGG {live['hash_agg']}): the table changed since the export, "
                             "re-run 'make export-marts'")
        finally:
            cur.close()
    finally:
        conn.close()


def run_check(out_dir: Path, offline: bool) -> int:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("FAIL  openpyxl is not installed. Run 'make setup' first.")
        return 2
    report = Report()
    print(f"Checking {rel(out_dir)}/ against {MANIFEST_NAME}" + (" (offline: files only)" if offline else ""))
    manifest = check_files(report, out_dir)
    if manifest is None:
        return 1
    if offline:
        print("INFO  Snowflake not queried (--offline): the live tables were not compared")
    else:
        try:
            check_live(report, manifest)
        except (sfconn.ConfigError, sfconn.ConnectError) as exc:
            print(f"FAIL  {exc}")
            return 2
        except ExportError as exc:
            print(f"FAIL  {exc}")
            return 1
        except ImportError:
            print("FAIL  snowflake-connector-python is not installed. Run 'make setup' first.")
            return 2
        except Exception as exc:
            from snowflake.connector.errors import Error as SnowflakeError

            if not isinstance(exc, SnowflakeError):
                raise
            print(f"FAIL  {sfconn.describe_error(exc)}")  # the connector's message, never credentials
            return 1
    extra = sorted(p.name for p in out_dir.glob("*.csv")
                   if p.name not in {t["csv"] for t in manifest["tables"]})
    if extra:
        report.warn("extra files", f"CSVs not in the manifest: {extra}")
    scope = "files" if offline else "files and live tables"
    if report.failures:
        print(f"FAIL  export check: {report.failures} check(s) failed ({scope}; exported "
              f"{manifest['exported_at_utc']})")
        return 1
    print(f"OK    export check: {len(manifest['tables'])} tables, {scope} match {MANIFEST_NAME} "
          f"(exported {manifest['exported_at_utc']}), {report.warnings} warning(s)")
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Export the dashboard's MART tables as the BI user (CSV, xlsx, manifest), or check an export.")
    parser.add_argument("--check", action="store_true",
                        help="verify exports/ against its manifest and the live tables instead of exporting")
    parser.add_argument("--offline", action="store_true",
                        help="with --check: verify the files only, without connecting to Snowflake")
    parser.add_argument("--out-dir", type=Path, default=DEFAULT_OUT_DIR,
                        help="export directory (default: exports/ in the repository)")
    args = parser.parse_args(argv)
    if args.offline and not args.check:
        parser.error("--offline only applies to --check")
    out_dir = args.out_dir if args.out_dir.is_absolute() else Path.cwd() / args.out_dir
    if args.check:
        return run_check(out_dir, args.offline)
    return run_export(out_dir)


if __name__ == "__main__":
    sys.exit(main())
